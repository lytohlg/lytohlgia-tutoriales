#!/usr/bin/env python3
"""Rellena la plantilla con los datos de clientes.csv: un .docx por cliente.

Regla del programa: NINGÚN documento sale si le queda un marcador o si le falta
un dato. Es preferible un contrato menos en la carpeta que un contrato con el
nombre del cliente anterior.

Uso:
    python3 rellenar.py                      # con clientes.csv
    python3 rellenar.py --csv clientes_con_hueco.csv
"""
import argparse
import csv
import os
import re
import sys
import unicodedata

from docx import Document

PLANTILLA = "plantilla/contrato_plantilla.docx"
SALIDA = "salida"
CAMPOS = ("cliente", "cif", "importe", "concepto", "fecha")
MARCAS = {
    "{{CLIENTE}}": "cliente",
    "{{CIF}}": "cif",
    "{{IMPORTE}}": "importe",
    "{{CONCEPTO}}": "concepto",
    "{{FECHA}}": "fecha",
}


def slug(texto):
    plano = "".join(c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "_", plano.lower()).strip("_")


def sustituir_parrafo(p, datos):
    """Sustituye sobre el TEXTO DEL PÁRRAFO, no run a run.

    Word parte un marcador en varios trozos cuando alguien edita encima de la
    frase: ahí "{{CLIENTE}}" puede estar en dos o tres runs. Mirando el párrafo
    entero, el marcador vuelve a aparecer completo. El precio: dentro de ese
    párrafo se pierde el formato variado (se queda el del primer trozo).
    """
    texto = p.text
    nuevo = texto
    for marca, campo in MARCAS.items():
        nuevo = nuevo.replace(marca, datos[campo])
    if nuevo == texto:
        return False
    runs = list(p.runs)
    for r in runs[1:]:
        r._element.getparent().remove(r._element)
    if runs:
        runs[0].text = nuevo
    else:
        p.add_run(nuevo)
    return True


def texto_documento(doc):
    partes = [p.text for p in doc.paragraphs]
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                partes.extend(p.text for p in celda.paragraphs)
    return "\n".join(partes)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", default="clientes.csv")
    args = ap.parse_args()

    if not os.path.exists(PLANTILLA):
        print("falta la plantilla:", PLANTILLA)
        return 2
    if not os.path.exists(args.csv):
        print("falta el fichero de datos:", args.csv)
        return 2

    with open(args.csv, encoding="utf-8", newline="") as f:
        filas = list(csv.DictReader(f))
    os.makedirs(SALIDA, exist_ok=True)

    generados = []
    detenidos = []
    for n, fila in enumerate(filas, 1):
        faltan = [c for c in CAMPOS if not (fila.get(c) or "").strip()]
        if faltan:
            nombre = (fila.get("cliente") or "").strip() or "sin nombre"
            print("%2d  %-26s DETENIDO: falta %s -> no se emite" % (n, nombre, ", ".join(faltan)))
            detenidos.append(n)
            continue

        datos = {c: fila[c].strip() for c in CAMPOS}
        doc = Document(PLANTILLA)
        for p in doc.paragraphs:
            sustituir_parrafo(p, datos)
        for tabla in doc.tables:
            for fila_t in tabla.rows:
                for celda in fila_t.cells:
                    for p in celda.paragraphs:
                        sustituir_parrafo(p, datos)

        restos = [m for m in MARCAS if m in texto_documento(doc)]
        if restos:
            print("%2d  %-26s DETENIDO: quedan marcadores %s" % (n, datos["cliente"], ", ".join(restos)))
            detenidos.append(n)
            continue

        destino = os.path.join(SALIDA, "contrato_%02d_%s.docx" % (n, slug(datos["cliente"])))
        doc.save(destino)
        generados.append(destino)
        print("%2d  %-26s %-28s -> %s (%d bytes)"
              % (n, datos["cliente"], datos["importe"] + " EUR", os.path.basename(destino),
                 os.path.getsize(destino)))

    print()
    print("RESUMEN: %d documentos generados, %d detenido%s, %d filas leídas"
          % (len(generados), len(detenidos), "" if len(detenidos) == 1 else "s", len(filas)))
    if detenidos:
        print("No se ha emitido ningún contrato a medias: los detenidos no tienen fichero.")
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
