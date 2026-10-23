#!/usr/bin/env python3
"""Los límites, medidos aquí y ahora (no "leídos en la documentación").

1) Un .doc antiguo (formato binario) NO se abre.
2) python-docx no exporta a PDF y en esta imagen no hay conversor.
3) Un campo de Word (DATE) no se recalcula: viaja con el valor guardado.
"""
import importlib.metadata as md
import os
import platform
import shutil

from docx import Document
from docx.oxml.ns import qn

print("python-docx", md.version("python-docx"), "| Python", platform.python_version())
print()

# 1. .doc binario antiguo
ruta_doc = "plantilla/contrato_antiguo.doc"
os.makedirs("plantilla", exist_ok=True)
with open(ruta_doc, "wb") as f:
    f.write(bytes.fromhex("d0cf11e0a1b11ae1") + b"\x00" * 512)
try:
    Document(ruta_doc)
    print("1. .doc antiguo: SE ABRIO (inesperado)")
except Exception as e:
    print("1. .doc antiguo: NO se abre ->", type(e).__name__)
    print("   lo que sí lee python-docx es .docx (el paquete OOXML), no el binario de 1997")

# 2. PDF
conversor = shutil.which("libreoffice") or shutil.which("soffice")
print("2. exportar a PDF: libreoffice/soffice ->", conversor or "no instalado")
print("   python-docx no convierte: si necesitas PDF, hace falta otro paso (y así se dice en el video)")

# 3. campo de Word
doc = Document("plantilla/contrato_plantilla.docx")
campos = doc.element.body.findall(".//" + qn("w:fldSimple"))
if not campos:
    print("3. campos de Word: ninguno en la plantilla (inesperado)")
for campo in campos:
    instruccion = (campo.get(qn("w:instr")) or "").strip()
    guardado = "".join(t.text or "" for t in campo.findall(".//" + qn("w:t")))
    print("3. campo de Word %r -> valor guardado %r" % (instruccion, guardado))
    print("   no se recalcula al rellenar: eso lo hace Word al abrir el documento")

# 4. La negrita DENTRO de un párrafo que lleva marcador.
import rellenar  # noqa: E402  (está al lado: es el mismo programa que rellena)

prueba = Document()
p = prueba.add_paragraph()
p.add_run("El cliente ")
en_negrita = p.add_run("{{CLIENTE}}")
en_negrita.bold = True
p.add_run(" firma abajo.")
antes = ["negrita" if r.bold else "normal" for r in p.runs]
rellenar.sustituir_parrafo(p, {"cliente": "ACME SL", "cif": "B1", "importe": "1",
                               "concepto": "c", "fecha": "01/01/2026"})
despues = ["negrita" if r.bold else "normal" for r in p.runs]
print("4. párrafo con el marcador en negrita: %d trozos %s -> %d trozo(s) %s"
      % (len(antes), antes, len(despues), despues))
print("   el marcador se sustituye, pero ese párrafo pierde la negrita: en la plantilla,")
print("   el marcador va en un párrafo de formato uniforme; la negrita se pone fuera de él")
