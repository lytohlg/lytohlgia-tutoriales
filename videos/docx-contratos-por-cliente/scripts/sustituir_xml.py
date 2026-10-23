#!/usr/bin/env python3
"""La forma que parece más rápida: abrir el .docx como ZIP y cambiar el XML.

Funciona mientras el marcador sea una sola pieza de texto dentro del XML.
En cuanto Word lo ha partido en dos trozos, la cadena "{{CLIENTE}}" NO existe
en el XML y el reemplazo no encuentra nada. El documento sale, sí: con el
marcador dentro.
"""
import argparse
import csv
import os
import zipfile

MARCAS = {
    "{{CLIENTE}}": "cliente",
    "{{CIF}}": "cif",
    "{{IMPORTE}}": "importe",
    "{{CONCEPTO}}": "concepto",
    "{{FECHA}}": "fecha",
}
PLANTILLA = "plantilla/contrato_plantilla.docx"
SALIDA = "salida_xml"


def sustituir_xml(origen, destino, datos):
    with zipfile.ZipFile(origen) as z:
        piezas = {n: z.read(n) for n in z.namelist()}
    xml = piezas["word/document.xml"].decode("utf-8")
    for marca, campo in MARCAS.items():
        xml = xml.replace(marca, datos[campo])
    piezas["word/document.xml"] = xml.encode("utf-8")
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as z:
        for nombre, datos_pieza in piezas.items():
            z.writestr(nombre, datos_pieza)
    return {m for m in MARCAS if m in xml}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", default="clientes.csv")
    args = ap.parse_args()
    with open(args.csv, encoding="utf-8", newline="") as f:
        filas = [x for x in csv.DictReader(f) if (x.get("cliente") or "").strip()]
    os.makedirs(SALIDA, exist_ok=True)

    con_restos = 0
    restos_vistos = set()
    for n, fila in enumerate(filas, 1):
        datos = {c: (fila.get(c) or "").strip() for c in MARCAS.values()}
        destino = os.path.join(SALIDA, "contrato_%02d.docx" % n)
        restos = sustituir_xml(PLANTILLA, destino, datos)
        if restos:
            con_restos += 1
            restos_vistos |= restos
    print("documentos escritos: %d" % len(filas))
    print("lo que dice ESTE script: %d de %d documentos con marcadores sin sustituir"
          % (con_restos, len(filas)))
    if restos_vistos:
        print("marcadores que ha visto: %s" % ", ".join(sorted(restos_vistos)))
    print()
    # La prueba de verdad: releer los ficheros escritos con python-docx. El marcador
    # partido no existe como cadena en el XML, así que la búsqueda de arriba no lo ve.
    from docx import Document

    callados = 0
    for n, fila in enumerate(filas, 1):
        doc = Document(os.path.join(SALIDA, "contrato_%02d.docx" % n))
        textos = [p.text for p in doc.paragraphs]
        for tabla in doc.tables:
            for fila_t in tabla.rows:
                for celda in fila_t.cells:
                    textos.extend(p.text for p in celda.paragraphs)
        sobran = sorted({m for m in MARCAS if m in "\n".join(textos)})
        if sobran:
            callados += 1
            if n == 1:
                print("al releer el documento 1 aparece: %s" % ", ".join(sobran))
    print("lo que dice la RElectura: %d de %d documentos llevan un marcador dentro"
          % (callados, len(filas)))
    if callados:
        print("y nadie se habría enterado: el .docx se abre y se ve \"{{CLIENTE}}\" impreso")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
