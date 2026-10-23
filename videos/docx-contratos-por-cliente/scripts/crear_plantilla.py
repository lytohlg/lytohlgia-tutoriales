#!/usr/bin/env python3
"""Crea la plantilla de Word del ejemplo (marcadores + tabla de conceptos).

Todo es inventado. La plantilla reproduce DOS cosas que pasan en la vida real
y que luego hay que saber manejar:

  1) Un marcador PARTIDO en dos trozos: "{{CLI" + "ENTE}}". Es lo que deja Word
     cuando alguien edita encima de una frase; el marcador sigue leyéndose bien
     en pantalla, pero por dentro ya no es una sola pieza de texto.
  2) Un campo FECHA de Word (w:fldSimple). No es texto: es una instrucción que
     Word recalcula al abrir. python-docx no la recalcula.
"""
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

os.makedirs("plantilla", exist_ok=True)

doc = Document()
titulo = doc.add_heading("CONTRATO DE SERVICIOS PROFESIONALES", level=1)

# Párrafo con el marcador del cliente PARTIDO en dos trozos (como lo deja Word).
p = doc.add_paragraph()
p.add_run("Estimado/a ")
p.add_run("{{CLI")
p.add_run("ENTE}}")
p.add_run(": el despacho prestará el servicio contratado en las condiciones de abajo.")

# Marcador entero, en una sola pieza.
p = doc.add_paragraph()
p.add_run("Con CIF ")
p.add_run("{{CIF}}")

# Tabla de conceptos: dos marcadores dentro de celdas (no son párrafos del cuerpo).
tabla = doc.add_table(rows=3, cols=2)
tabla.style = "Table Grid"
tabla.cell(0, 0).text = "Concepto"
tabla.cell(0, 1).text = "Importe"
tabla.cell(1, 0).text = "{{CONCEPTO}}"
tabla.cell(1, 1).text = "{{IMPORTE}} €"
tabla.cell(2, 0).text = "Forma de pago"
tabla.cell(2, 1).text = "Domiciliación mensual"

p = doc.add_paragraph()
p.add_run("Fecha de firma: ")
p.add_run("{{FECHA}}")

# Campo FECHA de Word: instrucción, no texto. Aquí queda el valor guardado.
p = doc.add_paragraph("Fecha del sistema: ")
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), " DATE ")
run = OxmlElement("w:r")
texto = OxmlElement("w:t")
texto.text = "01/01/2026"
run.append(texto)
fld.append(run)
p._p.append(fld)

salida = os.path.join("plantilla", "contrato_plantilla.docx")
doc.save(salida)
print("plantilla escrita:", salida, os.path.getsize(salida), "bytes")
print("párrafos:", len(doc.paragraphs), "| tablas:", len(doc.tables))
